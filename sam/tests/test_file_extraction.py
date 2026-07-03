"""Tests d'extraction de marchandises depuis fichiers tabulaires / Office."""

from __future__ import annotations

import csv
import io
import unittest

from docx import Document
from openpyxl import Workbook

from sam.api import (
    _build_structured_merchandise_item,
    _extract_items_from_csv,
    _extract_items_from_docx,
    _extract_items_from_tabular_rows,
    _extract_items_from_xlsx,
    _resolve_upload_extension,
)


class TestUploadExtension(unittest.TestCase):
    def test_resolve_from_filename(self) -> None:
        self.assertEqual(_resolve_upload_extension("liste.xlsx", None), "xlsx")
        self.assertEqual(_resolve_upload_extension("note.docx", "text/plain"), "docx")

    def test_resolve_from_mime(self) -> None:
        self.assertEqual(
            _resolve_upload_extension(
                "export",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            "xlsx",
        )


class TestTabularExtraction(unittest.TestCase):
    def test_multi_column_header_builds_structured_item(self) -> None:
        rows = [
            ["Désignation", "Matière", "Usage", "Caractéristiques", "Qté"],
            ["Sac de voyage", "cuir", "transport", "neuf", "2"],
        ]
        _, items = _extract_items_from_tabular_rows(rows, max_items=10)
        self.assertEqual(len(items), 1)
        self.assertIn("Produit : Sac de voyage", items[0])
        self.assertIn("Composition", items[0])
        self.assertIn("Usage", items[0])

    def test_full_excel_header(self) -> None:
        rows = [
            [
                "Désignation",
                "Matière / composition",
                "Usage",
                "Caractéristiques",
                "Qté",
                "Unité",
                "Pays d'origine",
                "Valeur",
                "Devise",
            ],
            [
                "Téléviseur LED 55''",
                "plastique, verre",
                "domestique",
                "neuf",
                "10",
                "u",
                "Chine",
                "250000",
                "XOF",
            ],
        ]
        _, items = _extract_items_from_tabular_rows(rows, max_items=10)
        self.assertEqual(len(items), 1)
        self.assertIn("Produit : Téléviseur LED 55''", items[0])
        self.assertIn("Origine", items[0])
        self.assertIn("Chine", items[0])
        self.assertIn("Valeur", items[0])
        self.assertIn("250000 XOF", items[0])
        self.assertIn("10 u", items[0])

    def test_csv_semicolon(self) -> None:
        text = "Produit;Matière;Usage\nOrdinateur portable;plastique;bureau\n"
        _, items = _extract_items_from_csv(text, max_items=10)
        self.assertEqual(len(items), 1)
        self.assertIn("Ordinateur portable", items[0])


class TestExcelExtraction(unittest.TestCase):
    def test_xlsx_two_rows(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.append(["Produit", "Composition", "Usage"])
        ws.append(["Chaise de bureau", "acier", "bureau"])
        buf = io.BytesIO()
        wb.save(buf)
        _, items = _extract_items_from_xlsx(buf.getvalue(), max_items=10)
        self.assertEqual(len(items), 1)
        self.assertIn("Chaise de bureau", items[0])


class TestWordExtraction(unittest.TestCase):
    def test_docx_table(self) -> None:
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.rows[0].cells[0].text = "Produit"
        table.rows[0].cells[1].text = "Matière"
        table.rows[0].cells[2].text = "Usage"
        table.rows[1].cells[0].text = "Téléphone"
        table.rows[1].cells[1].text = "plastique"
        table.rows[1].cells[2].text = "communication"
        buf = io.BytesIO()
        doc.save(buf)
        _, items = _extract_items_from_docx(buf.getvalue(), max_items=10)
        self.assertEqual(len(items), 1)
        self.assertIn("Téléphone", items[0])


class TestStructuredItem(unittest.TestCase):
    def test_designation_only(self) -> None:
        self.assertEqual(
            _build_structured_merchandise_item("Câble USB"),
            "Câble USB",
        )


if __name__ == "__main__":
    unittest.main()
